from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass
from datetime import datetime
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session, joinedload

from ..config import get_settings
from ..models import InventoryItem, PurchaseImportItemMeta, PurchaseImportState, PurchaseOrder, PurchaseOrderItem, PurchaseOrderStatus
from ..schemas import (
    PurchaseImportItemRead,
    PurchaseImportedItemUpdate,
    PurchaseImportResponse,
    PurchaseImportStateRead,
    PurchasePendingReceiptDeleteResponse,
)
from .bootstrap import normalize_text, to_date, to_decimal


PURCHASE_SHEET_CONFIGS = {
    "Sheet1": {"header_row": 2, "material_key": "物料名称", "quantity_key": "需求数量 "},
    "佳时坤": {"header_row": 2, "material_key": "产品名称", "quantity_key": "数量 "},
}


@dataclass
class ParsedPurchaseRow:
    sheet_name: str
    source_row_number: int
    requester: str | None
    purchase_category: str | None
    project_name: str | None
    supplier_name: str | None
    contract_no: str | None
    requested_at: object
    ordered_at: object
    notes: str | None
    material_name: str
    specification: str | None
    requested_quantity: object
    unit: str | None
    unit_price: object | None
    tax_rate: object | None
    total_amount: object | None
    expected_arrival: object


def get_or_create_purchase_import_states(session: Session) -> list[PurchaseImportState]:
    states = {
        state.sheet_name: state
        for state in session.scalars(select(PurchaseImportState).order_by(PurchaseImportState.sheet_name.asc())).all()
    }

    created = False
    for sheet_name, config in PURCHASE_SHEET_CONFIGS.items():
        if sheet_name not in states:
            state = PurchaseImportState(
                sheet_name=sheet_name,
                header_row=config["header_row"],
                last_imported_row=config["header_row"],
            )
            session.add(state)
            states[sheet_name] = state
            created = True

    if created:
        session.commit()
        for state in states.values():
            session.refresh(state)

    return [states[sheet_name] for sheet_name in PURCHASE_SHEET_CONFIGS]


def serialize_purchase_import_state(state: PurchaseImportState) -> PurchaseImportStateRead:
    return PurchaseImportStateRead(
        sheet_name=state.sheet_name,
        header_row=state.header_row,
        last_imported_row=state.last_imported_row,
        last_workbook_name=state.last_workbook_name,
        updated_at=state.updated_at,
    )


def update_purchase_import_states(session: Session, updates: list[dict[str, int]]) -> list[PurchaseImportStateRead]:
    states = {state.sheet_name: state for state in get_or_create_purchase_import_states(session)}

    for update in updates:
        sheet_name = update["sheet_name"]
        if sheet_name not in states:
            raise ValueError(f"Unknown purchase sheet: {sheet_name}")

        state = states[sheet_name]
        next_row = update["last_imported_row"]
        if next_row < state.header_row:
            raise ValueError(f"{sheet_name} 的行号不能小于表头行 {state.header_row}")
        state.last_imported_row = next_row

    session.commit()
    return [serialize_purchase_import_state(states[sheet_name]) for sheet_name in PURCHASE_SHEET_CONFIGS]


def _match_inventory_item(session: Session, material_name: str, specification: str | None) -> InventoryItem | None:
    return session.scalar(
        select(InventoryItem)
        .where(
            InventoryItem.material_name == material_name,
            InventoryItem.specification == specification,
        )
        .order_by(InventoryItem.id.asc())
    )


def _serialize_imported_item(item: PurchaseOrderItem) -> PurchaseImportItemRead:
    return PurchaseImportItemRead(
        purchase_item_id=item.id,
        purchase_order_id=item.order_id,
        sheet_name=item.order.sheet_name,
        source_row_number=item.import_meta.source_row_number if item.import_meta else 0,
        requester=item.order.requester,
        purchase_category=item.order.purchase_category,
        project_name=item.order.project_name,
        supplier_name=item.order.supplier_name,
        contract_no=item.order.contract_no,
        material_name=item.material_name,
        specification=item.specification,
        requested_quantity=item.requested_quantity,
        unit=item.unit,
        expected_arrival=item.expected_arrival,
        inventory_item_id=item.inventory_item_id,
    )


def import_purchase_workbook_incremental(
    session: Session,
    workbook_bytes: bytes,
    *,
    workbook_name: str,
    persist_uploaded_copy: bool = True,
) -> PurchaseImportResponse:
    settings = get_settings()
    workbook = load_workbook(BytesIO(workbook_bytes), data_only=True)
    states = {state.sheet_name: state for state in get_or_create_purchase_import_states(session)}
    warnings: list[str] = []
    parsed_rows: list[ParsedPurchaseRow] = []
    max_rows_by_sheet: dict[str, int] = {}

    for sheet_name, config in PURCHASE_SHEET_CONFIGS.items():
        if sheet_name not in workbook.sheetnames:
            warnings.append(f"上传文件缺少工作表 {sheet_name}，本次已跳过。")
            continue

        ws = workbook[sheet_name]
        header_row = config["header_row"]
        state = states[sheet_name]
        start_row = max(state.last_imported_row + 1, header_row + 1)
        headers = [cell.value for cell in ws[header_row]]
        max_rows_by_sheet[sheet_name] = state.last_imported_row

        for row_number in range(start_row, ws.max_row + 1):
            row = [cell.value for cell in ws[row_number]]
            row_data = dict(zip(headers, row))
            material_name = normalize_text(row_data.get(config["material_key"]))
            if not material_name:
                continue

            parsed_rows.append(
                ParsedPurchaseRow(
                    sheet_name=sheet_name,
                    source_row_number=row_number,
                    requester=normalize_text(row_data.get("请购人")),
                    purchase_category=normalize_text(row_data.get("请购类别")),
                    project_name=normalize_text(row_data.get("项目")),
                    supplier_name=normalize_text(row_data.get("供应商名称")),
                    contract_no=normalize_text(row_data.get("合同编号")),
                    requested_at=to_date(row_data.get("需求日期")),
                    ordered_at=to_date(row_data.get("合同日期")),
                    notes=normalize_text(row_data.get("交易方")),
                    material_name=material_name,
                    specification=normalize_text(row_data.get("规格型号")),
                    requested_quantity=to_decimal(row_data.get(config["quantity_key"])),
                    unit=normalize_text(row_data.get("单位")),
                    unit_price=to_decimal(row_data.get("含税单价")) if row_data.get("含税单价") not in (None, "") else None,
                    tax_rate=to_decimal(row_data.get("税率")) if row_data.get("税率") not in (None, "") else None,
                    total_amount=to_decimal(row_data.get("总价")) if row_data.get("总价") not in (None, "") else None,
                    expected_arrival=to_date(row_data.get("到货日期")),
                )
            )
            max_rows_by_sheet[sheet_name] = row_number

    if not parsed_rows:
        for sheet_name, state in states.items():
            state.last_workbook_name = workbook_name
        session.commit()
        if persist_uploaded_copy:
            archive_uploaded_purchase_workbook(workbook_bytes, workbook_name)
        return PurchaseImportResponse(
            workbook_name=workbook_name,
            imported_order_count=0,
            imported_item_count=0,
            unmatched_item_count=0,
            updated_states=[serialize_purchase_import_state(states[sheet_name]) for sheet_name in PURCHASE_SHEET_CONFIGS],
            imported_items=[],
            warnings=warnings,
        )

    parsed_rows.sort(
        key=lambda row: (
            row.sheet_name,
            row.supplier_name or "",
            row.requester or "",
            row.purchase_category or "",
            row.project_name or "",
            row.contract_no or "",
            row.source_row_number,
        )
    )

    grouped_rows: dict[tuple[str, str | None, str | None, str | None, str | None, str | None], list[ParsedPurchaseRow]] = defaultdict(list)
    for row in parsed_rows:
        key = (
            row.sheet_name,
            row.supplier_name,
            row.requester,
            row.purchase_category,
            row.project_name,
            row.contract_no,
        )
        grouped_rows[key].append(row)

    created_items: list[PurchaseOrderItem] = []
    created_order_count = 0
    unmatched_item_count = 0

    for rows in grouped_rows.values():
        first = rows[0]
        order = PurchaseOrder(
            sheet_name=first.sheet_name,
            requester=first.requester,
            purchase_category=first.purchase_category,
            project_name=first.project_name,
            supplier_name=first.supplier_name,
            status=PurchaseOrderStatus.pending,
            requested_at=first.requested_at,
            ordered_at=first.ordered_at,
            contract_no=first.contract_no,
            notes=first.notes,
        )
        session.add(order)
        session.flush()
        created_order_count += 1

        for line_no, row in enumerate(rows, start=1):
            inventory_item = _match_inventory_item(session, row.material_name, row.specification)
            if not inventory_item:
                unmatched_item_count += 1

            item = PurchaseOrderItem(
                order_id=order.id,
                inventory_item_id=inventory_item.id if inventory_item else None,
                line_no=line_no,
                material_name=row.material_name,
                specification=row.specification,
                requested_quantity=row.requested_quantity,
                received_quantity=0,
                unit=row.unit,
                unit_price=row.unit_price,
                tax_rate=row.tax_rate,
                total_amount=row.total_amount,
                expected_arrival=row.expected_arrival,
            )
            session.add(item)
            session.flush()
            session.add(
                PurchaseImportItemMeta(
                    purchase_item_id=item.id,
                    source_row_number=row.source_row_number,
                )
            )
            created_items.append(item)

    for sheet_name, last_row in max_rows_by_sheet.items():
        state = states[sheet_name]
        state.last_imported_row = max(state.last_imported_row, last_row)
        state.last_workbook_name = workbook_name

    session.commit()

    created_item_ids = [item.id for item in created_items]
    refreshed_items = session.scalars(
        select(PurchaseOrderItem)
        .options(joinedload(PurchaseOrderItem.order), joinedload(PurchaseOrderItem.import_meta))
        .where(PurchaseOrderItem.id.in_(created_item_ids))
        .order_by(PurchaseOrderItem.id.asc())
    ).unique().all()

    if persist_uploaded_copy:
        archive_uploaded_purchase_workbook(workbook_bytes, workbook_name)

    return PurchaseImportResponse(
        workbook_name=workbook_name,
        imported_order_count=created_order_count,
        imported_item_count=len(refreshed_items),
        unmatched_item_count=unmatched_item_count,
        updated_states=[serialize_purchase_import_state(states[sheet_name]) for sheet_name in PURCHASE_SHEET_CONFIGS],
        imported_items=[_serialize_imported_item(item) for item in refreshed_items],
        warnings=warnings,
    )


def archive_uploaded_purchase_workbook(workbook_bytes: bytes, workbook_name: str) -> Path:
    settings = get_settings()
    archive_dir = settings.purchase_workbook.parent / "uploaded"
    archive_dir.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    archive_path = archive_dir / f"{timestamp}_{Path(workbook_name).name}"
    archive_path.write_bytes(workbook_bytes)
    return archive_path


def update_imported_purchase_items(session: Session, updates: list[PurchaseImportedItemUpdate]) -> list[PurchaseImportItemRead]:
    updated_ids: list[int] = []
    for update in updates:
        item = session.get(PurchaseOrderItem, update.purchase_item_id)
        if not item:
            raise ValueError(f"采购明细 {update.purchase_item_id} 不存在")

        item.material_name = update.material_name.strip()
        item.specification = normalize_text(update.specification)
        item.requested_quantity = update.requested_quantity
        item.unit = normalize_text(update.unit)
        item.expected_arrival = update.expected_arrival

        if update.inventory_item_id is None:
            item.inventory_item_id = None
        else:
            inventory_item = session.get(InventoryItem, update.inventory_item_id)
            if not inventory_item:
                raise ValueError(f"库存项目 {update.inventory_item_id} 不存在")
            item.inventory_item_id = inventory_item.id

        updated_ids.append(item.id)

    session.commit()
    refreshed_items = session.scalars(
        select(PurchaseOrderItem)
        .options(joinedload(PurchaseOrderItem.order), joinedload(PurchaseOrderItem.import_meta))
        .where(PurchaseOrderItem.id.in_(updated_ids))
        .order_by(PurchaseOrderItem.id.asc())
    ).unique().all()
    return [_serialize_imported_item(item) for item in refreshed_items]


def delete_pending_purchase_items(session: Session, item_ids: list[int]) -> list[int]:
    normalized_ids = sorted(set(item_ids))
    if not normalized_ids:
        raise ValueError("No purchase items selected for deletion.")

    items = session.scalars(
        select(PurchaseOrderItem)
        .options(joinedload(PurchaseOrderItem.order))
        .where(PurchaseOrderItem.id.in_(normalized_ids))
    ).unique().all()
    if not items:
        raise ValueError("Selected purchase items were not found.")

    deleted_ids: list[int] = []
    touched_orders: set[int] = set()
    for item in items:
        if (item.received_quantity or 0) > 0:
            raise ValueError(f"采购明细 {item.id} 已收货，不能删除。")

        touched_orders.add(item.order_id)
        session.delete(item)
        deleted_ids.append(item.id)

    session.flush()

    for order_id in touched_orders:
        remaining_items = session.scalars(
            select(PurchaseOrderItem.id).where(PurchaseOrderItem.order_id == order_id)
        ).all()
        if not remaining_items:
            order = session.get(PurchaseOrder, order_id)
            if order:
                session.delete(order)

    session.commit()
    return deleted_ids
