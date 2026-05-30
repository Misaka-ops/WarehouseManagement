from __future__ import annotations

from io import BytesIO
from datetime import date, datetime
from decimal import Decimal
from itertools import groupby
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import delete, func, select, text, update
from sqlalchemy.orm import Session, joinedload

from ..config import get_settings
from ..database import Base, engine
from ..models import (
    InventoryItem,
    InventoryTransaction,
    Location,
    PurchaseOrder,
    PurchaseOrderItem,
    PurchaseOrderStatus,
    Supplier,
    TransactionType,
)
from ..schemas import InventoryImportResponse


def normalize_text(value) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def to_decimal(value) -> Decimal:
    if value in (None, ""):
        return Decimal("0")
    return Decimal(str(value))


def to_date(value) -> date | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    if not text:
        return None
    if "2026年之前" in text:
        return None
    for fmt in ("%Y-%m-%d", "%Y/%m/%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def get_or_create_supplier(session: Session, name: str | None) -> Supplier | None:
    if not name:
        return None
    supplier = session.scalar(select(Supplier).where(Supplier.name == name))
    if supplier:
        return supplier
    supplier = Supplier(name=name)
    session.add(supplier)
    session.flush()
    return supplier


def get_or_create_location(session: Session, name: str | None) -> Location | None:
    if not name:
        return None
    location = session.scalar(select(Location).where(Location.name == name))
    if location:
        return location
    location = Location(name=name)
    session.add(location)
    session.flush()
    return location


def create_schema() -> None:
    Base.metadata.create_all(bind=engine)
    migrate_inventory_item_unique_constraint()


def _inventory_item_unique_indexes(connection) -> list[list[str]]:
    try:
        rows = connection.execute(text("PRAGMA index_list('inventory_items')")).fetchall()
    except Exception:
        return []

    index_columns: list[list[str]] = []
    for row in rows:
        if len(row) < 3 or not row[2]:
            continue
        index_name = row[1]
        column_rows = connection.execute(text(f"PRAGMA index_info('{index_name}')")).fetchall()
        columns = [column_row[2] for column_row in column_rows if len(column_row) >= 3 and column_row[2]]
        if columns:
            index_columns.append(columns)
    return index_columns


def migrate_inventory_item_unique_constraint() -> None:
    if engine.dialect.name != "sqlite":
        return

    with engine.begin() as connection:
        unique_indexes = _inventory_item_unique_indexes(connection)
        if not unique_indexes:
            return

        target_columns = ["material_name", "specification", "unit", "supplier_id", "location_id"]
        legacy_columns = ["material_name", "specification", "supplier_id", "location_id"]

        has_target_constraint = any(columns == target_columns for columns in unique_indexes)
        has_legacy_constraint = any(columns == legacy_columns for columns in unique_indexes)
        if has_target_constraint or not has_legacy_constraint:
            return

        connection.execute(text("PRAGMA foreign_keys=OFF"))
        try:
            connection.execute(
                text(
                    """
                    CREATE TABLE inventory_items__new (
                        id INTEGER NOT NULL PRIMARY KEY,
                        requester VARCHAR(80),
                        purchase_category VARCHAR(80),
                        project_name VARCHAR(120),
                        material_name VARCHAR(200) NOT NULL,
                        specification VARCHAR(255),
                        unit VARCHAR(40),
                        supplier_id INTEGER REFERENCES suppliers (id),
                        location_id INTEGER REFERENCES locations (id),
                        quantity_on_hand NUMERIC(12, 2) NOT NULL,
                        notes TEXT,
                        last_receipt_at DATE,
                        last_issue_at DATE,
                        created_at DATETIME NOT NULL,
                        updated_at DATETIME NOT NULL,
                        UNIQUE (material_name, specification, unit, supplier_id, location_id)
                    )
                    """
                )
            )
            connection.execute(
                text(
                    """
                    INSERT INTO inventory_items__new (
                        id,
                        requester,
                        purchase_category,
                        project_name,
                        material_name,
                        specification,
                        unit,
                        supplier_id,
                        location_id,
                        quantity_on_hand,
                        notes,
                        last_receipt_at,
                        last_issue_at,
                        created_at,
                        updated_at
                    )
                    SELECT
                        id,
                        requester,
                        purchase_category,
                        project_name,
                        material_name,
                        specification,
                        unit,
                        supplier_id,
                        location_id,
                        quantity_on_hand,
                        notes,
                        last_receipt_at,
                        last_issue_at,
                        created_at,
                        updated_at
                    FROM inventory_items
                    """
                )
            )
            connection.execute(text("DROP TABLE inventory_items"))
            connection.execute(text("ALTER TABLE inventory_items__new RENAME TO inventory_items"))
            connection.execute(text("CREATE INDEX ix_inventory_items_material_name ON inventory_items (material_name)"))
        finally:
            connection.execute(text("PRAGMA foreign_keys=ON"))


def import_warehouse_workbook(session: Session, workbook_source: Path | BytesIO) -> tuple[int, int]:
    wb = load_workbook(workbook_source, data_only=True)
    ws = wb["原物料库存清单"]
    imported_item_count = 0
    imported_transaction_count = 0

    for row in ws.iter_rows(min_row=2, values_only=True):
        (
            requester,
            category,
            project,
            material,
            specification,
            _,
            unit,
            supplier_name,
            stock_qty,
            notes,
            location_name,
            receipt_date,
            receipt_qty,
            issue_qty_1,
            issue_date_1,
            issue_operator_1,
            issue_qty_2,
            issue_date_2,
            issue_operator_2,
            issue_qty_3,
            issue_date_3,
            issue_operator_3,
            *_,
        ) = row
        material_name = normalize_text(material)
        if not material_name:
            continue

        supplier = get_or_create_supplier(session, normalize_text(supplier_name))
        location = get_or_create_location(session, normalize_text(location_name))
        item = InventoryItem(
            requester=normalize_text(requester),
            purchase_category=normalize_text(category),
            project_name=normalize_text(project),
            material_name=material_name,
            specification=normalize_text(specification),
            unit=normalize_text(unit),
            supplier=supplier,
            location=location,
            quantity_on_hand=to_decimal(stock_qty),
            notes=normalize_text(notes),
            last_receipt_at=to_date(receipt_date),
            last_issue_at=max(
                (d for d in (to_date(issue_date_1), to_date(issue_date_2), to_date(issue_date_3)) if d is not None),
                default=None,
            ),
        )
        session.add(item)
        session.flush()
        imported_item_count += 1

        receipt_qty_decimal = to_decimal(receipt_qty)
        if receipt_qty_decimal > 0:
            session.add(
                InventoryTransaction(
                    item_id=item.id,
                    transaction_type=TransactionType.receipt,
                    quantity=receipt_qty_decimal,
                    occurred_on=to_date(receipt_date) or date.today(),
                    operator_name="系统导入",
                    reference_code="legacy-warehouse-import",
                    notes="由仓库 Excel 初始化导入",
                )
            )
            imported_transaction_count += 1

        issue_slots = [
            (issue_qty_1, issue_date_1, issue_operator_1, 1),
            (issue_qty_2, issue_date_2, issue_operator_2, 2),
            (issue_qty_3, issue_date_3, issue_operator_3, 3),
        ]
        for issue_qty, issue_date, issue_operator, slot_no in issue_slots:
            issue_qty_decimal = to_decimal(issue_qty)
            issue_occurred_on = to_date(issue_date)
            issue_operator_name = normalize_text(issue_operator)
            if issue_qty_decimal > 0:
                session.add(
                    InventoryTransaction(
                        item_id=item.id,
                        transaction_type=TransactionType.issue,
                        quantity=issue_qty_decimal,
                        occurred_on=issue_occurred_on or date.today(),
                        operator_name=issue_operator_name,
                        reference_code=f"legacy-warehouse-import-issue-{slot_no}",
                        notes="由仓库 Excel 初始化导入",
                    )
                )
                imported_transaction_count += 1

    return imported_item_count, imported_transaction_count


def relink_purchase_items(session: Session) -> int:
    relinked_count = 0
    purchase_items = session.scalars(select(PurchaseOrderItem).order_by(PurchaseOrderItem.id.asc())).all()
    for purchase_item in purchase_items:
        inventory_item = session.scalar(
            select(InventoryItem)
            .where(
                InventoryItem.material_name == purchase_item.material_name,
                InventoryItem.specification == purchase_item.specification,
            )
            .order_by(InventoryItem.id.asc())
        )
        purchase_item.inventory_item_id = inventory_item.id if inventory_item else None
        if inventory_item:
            relinked_count += 1

    return relinked_count


def replace_inventory_from_warehouse_workbook(
    session: Session,
    workbook_bytes: bytes,
    *,
    workbook_name: str,
    persist_uploaded_copy: bool = True,
) -> InventoryImportResponse:
    settings = get_settings()
    replaced_item_count = session.scalar(select(func.count()).select_from(InventoryItem)) or 0
    workbook_preview = load_workbook(BytesIO(workbook_bytes), data_only=True)
    if "原物料库存清单" not in workbook_preview.sheetnames:
        raise KeyError("原物料库存清单")

    session.execute(update(PurchaseOrderItem).values(inventory_item_id=None))
    session.execute(delete(InventoryTransaction))
    session.execute(delete(InventoryItem))
    session.execute(delete(Supplier))
    session.execute(delete(Location))
    session.flush()

    imported_item_count, imported_transaction_count = import_warehouse_workbook(session, BytesIO(workbook_bytes))
    relinked_purchase_item_count = relink_purchase_items(session)
    session.commit()

    if persist_uploaded_copy:
        settings.warehouse_workbook.parent.mkdir(parents=True, exist_ok=True)
        settings.warehouse_workbook.write_bytes(workbook_bytes)

    return InventoryImportResponse(
        workbook_name=workbook_name,
        imported_item_count=imported_item_count,
        imported_transaction_count=imported_transaction_count,
        relinked_purchase_item_count=relinked_purchase_item_count,
        replaced_item_count=int(replaced_item_count),
    )


def _format_excel_date(value: date | None) -> str | None:
    if value is None:
        return None
    return value.strftime("%Y-%m-%d")


def export_warehouse_workbook(session: Session) -> tuple[str, bytes]:
    settings = get_settings()
    workbook = load_workbook(settings.warehouse_workbook)
    if "原物料库存清单" not in workbook.sheetnames:
        raise KeyError("原物料库存清单")

    ws = workbook["原物料库存清单"]
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)

    items = session.scalars(
        select(InventoryItem)
        .options(
            joinedload(InventoryItem.supplier),
            joinedload(InventoryItem.location),
            joinedload(InventoryItem.transactions),
        )
        .order_by(InventoryItem.id.asc())
    ).unique().all()

    for row_index, item in enumerate(items, start=2):
        receipts = sorted(
            [tx for tx in item.transactions if tx.transaction_type == TransactionType.receipt],
            key=lambda tx: (tx.occurred_on, tx.id),
            reverse=True,
        )
        issues = sorted(
            [tx for tx in item.transactions if tx.transaction_type == TransactionType.issue],
            key=lambda tx: (tx.occurred_on, tx.id),
        )[:3]
        latest_receipt = receipts[0] if receipts else None

        row_values = [
            item.requester,
            item.purchase_category,
            item.project_name,
            item.material_name,
            item.specification,
            None,
            item.unit,
            item.supplier.name if item.supplier else None,
            float(item.quantity_on_hand),
            item.notes,
            item.location.name if item.location else None,
            _format_excel_date(latest_receipt.occurred_on if latest_receipt else item.last_receipt_at),
            float(latest_receipt.quantity) if latest_receipt else None,
        ]

        for slot in range(3):
            issue = issues[slot] if slot < len(issues) else None
            row_values.extend(
                [
                    float(issue.quantity) if issue else None,
                    _format_excel_date(issue.occurred_on) if issue else None,
                    issue.operator_name if issue else None,
                ]
            )

        for column_index, value in enumerate(row_values, start=1):
            ws.cell(row=row_index, column=column_index, value=value)

    output = BytesIO()
    workbook.save(output)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    export_name = f"{settings.warehouse_workbook.stem}_导出_{timestamp}{settings.warehouse_workbook.suffix}"
    return export_name, output.getvalue()


def import_purchase_workbook(session: Session, workbook_path: Path) -> None:
    wb = load_workbook(workbook_path, data_only=True)
    sheet_configs = {"Sheet1": {"header_row": 2, "material_key": "物料名称", "quantity_key": "需求数量 "}, "佳时坤": {"header_row": 2, "material_key": "产品名称", "quantity_key": "数量 "}}

    for sheet_name, config in sheet_configs.items():
        ws = wb[sheet_name]
        headers = [cell.value for cell in ws[config["header_row"]]]
        rows = []
        for row in ws.iter_rows(min_row=config["header_row"] + 1, values_only=True):
            row_data = dict(zip(headers, row))
            material_name = normalize_text(row_data.get(config["material_key"]))
            if not material_name:
                continue
            rows.append(row_data)

        rows.sort(key=lambda r: (
            normalize_text(r.get("供应商名称")) or "",
            normalize_text(r.get("请购人")) or "",
            normalize_text(r.get("请购类别")) or "",
            normalize_text(r.get("项目")) or "",
            normalize_text(r.get("合同编号")) or "",
        ))

        for _, grouped in groupby(
            rows,
            key=lambda r: (
                normalize_text(r.get("供应商名称")),
                normalize_text(r.get("请购人")),
                normalize_text(r.get("请购类别")),
                normalize_text(r.get("项目")),
                normalize_text(r.get("合同编号")),
            ),
        ):
            group = list(grouped)
            first = group[0]
            order = PurchaseOrder(
                sheet_name=sheet_name,
                requester=normalize_text(first.get("请购人")),
                purchase_category=normalize_text(first.get("请购类别")),
                project_name=normalize_text(first.get("项目")),
                supplier_name=normalize_text(first.get("供应商名称")),
                status=PurchaseOrderStatus.pending,
                requested_at=to_date(first.get("需求日期")),
                ordered_at=to_date(first.get("合同日期")),
                contract_no=normalize_text(first.get("合同编号")),
                notes=normalize_text(first.get("交易方")),
            )
            session.add(order)
            session.flush()

            for line_no, row in enumerate(group, start=1):
                supplier_name = normalize_text(row.get("供应商名称"))
                material_name = normalize_text(row.get(config["material_key"]))
                specification = normalize_text(row.get("规格型号"))
                inventory_item = session.scalar(
                    select(InventoryItem).where(
                        InventoryItem.material_name == material_name,
                        InventoryItem.specification == specification,
                    )
                )
                session.add(
                    PurchaseOrderItem(
                        order_id=order.id,
                        inventory_item_id=inventory_item.id if inventory_item else None,
                        line_no=line_no,
                        material_name=material_name,
                        specification=specification,
                        requested_quantity=to_decimal(row.get(config["quantity_key"])),
                        received_quantity=Decimal("0"),
                        unit=normalize_text(row.get("单位")),
                        unit_price=to_decimal(row.get("含税单价")) if row.get("含税单价") not in (None, "") else None,
                        tax_rate=to_decimal(row.get("税率")) if row.get("税率") not in (None, "") else None,
                        total_amount=to_decimal(row.get("总价")) if row.get("总价") not in (None, "") else None,
                        expected_arrival=to_date(row.get("到货日期")),
                    )
                )
