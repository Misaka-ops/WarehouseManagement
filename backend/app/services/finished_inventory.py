from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel
from sqlalchemy import and_, select
from sqlalchemy.orm import Session

from ..config import get_settings
from ..models import (
    FinishedInventoryItem as FinishedInventoryItemModel,
    FinishedInventoryTransaction as FinishedInventoryTransactionModel,
    TransactionType,
)
from ..schemas import (
    FinishedInventoryDashboardResponse,
    FinishedInventoryItemRead,
    FinishedInventorySummary,
    FinishedInventoryTransactionRead,
)
from .bootstrap import normalize_text, to_date


FINISHED_SHEET_NAME = "成品库存清单"
LOW_STOCK_THRESHOLD = Decimal("5")


class DuplicateFinishedInventoryItemError(ValueError):
    def __init__(self, row_id: int):
        super().__init__(f"Finished inventory item already exists at row {row_id}.")
        self.row_id = row_id


def _db_row_id(item_id: int) -> int:
    return -item_id


def _db_item_id(row_id: int) -> int | None:
    return abs(row_id) if row_id < 0 else None


@dataclass(frozen=True)
class FinishedInventoryRow:
    row_id: int
    material_name: str
    specification: str | None
    work_order_no: str | None
    quantity_on_hand: Decimal
    unit: str | None
    location_name: str | None
    receipt_date: date | None
    receipt_quantity: Decimal
    notes: str | None
    issue_quantity_1: Decimal
    issue_date_1: date | None
    issue_operator_1: str | None
    issue_quantity_2: Decimal
    issue_date_2: date | None
    issue_operator_2: str | None
    issue_quantity_3: Decimal
    issue_date_3: date | None
    issue_operator_3: str | None
    project_code: str | None
    producer_name: str | None
    customer_name: str | None


def _to_decimal(value) -> Decimal:
    if value in (None, ""):
        return Decimal("0")

    try:
        return Decimal(str(value).strip())
    except (InvalidOperation, AttributeError, ValueError):
        return Decimal("0")


def _to_finished_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float, Decimal)):
        try:
            converted = from_excel(float(value))
        except (TypeError, ValueError, OverflowError):
            return None
        if isinstance(converted, datetime):
            return converted.date()
        if isinstance(converted, date):
            return converted
        return None
    return to_date(value)


def _same_finished_item(
    row: FinishedInventoryRow,
    *,
    material_name: str,
    specification: str | None,
    unit: str | None,
    location_name: str | None,
) -> bool:
    return (
        row.material_name == material_name
        and row.specification == specification
        and row.unit == unit
        and row.location_name == location_name
    )


def _finished_match_conditions(
    *,
    material_name: str,
    specification: str | None,
    unit: str | None,
    location_name: str | None,
):
    conditions = [FinishedInventoryItemModel.material_name == material_name]
    conditions.append(
        FinishedInventoryItemModel.specification.is_(None)
        if specification is None
        else FinishedInventoryItemModel.specification == specification
    )
    conditions.append(FinishedInventoryItemModel.unit.is_(None) if unit is None else FinishedInventoryItemModel.unit == unit)
    conditions.append(
        FinishedInventoryItemModel.location_name.is_(None)
        if location_name is None
        else FinishedInventoryItemModel.location_name == location_name
    )
    return conditions


def _load_finished_rows() -> list[FinishedInventoryRow]:
    settings = get_settings()
    workbook_path = settings.warehouse_workbook
    workbook = load_workbook(workbook_path, data_only=True, read_only=True)

    if FINISHED_SHEET_NAME not in workbook.sheetnames:
        raise KeyError(FINISHED_SHEET_NAME)

    sheet = workbook[FINISHED_SHEET_NAME]
    rows: list[FinishedInventoryRow] = []

    for row_id, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        material_name = normalize_text(row[1] if len(row) > 1 else None)
        if not material_name:
            continue

        rows.append(
            FinishedInventoryRow(
                row_id=row_id,
                material_name=material_name,
                specification=normalize_text(row[2] if len(row) > 2 else None),
                work_order_no=normalize_text(row[3] if len(row) > 3 else None),
                quantity_on_hand=_to_decimal(row[4] if len(row) > 4 else None),
                unit=normalize_text(row[5] if len(row) > 5 else None),
                location_name=normalize_text(row[6] if len(row) > 6 else None),
                receipt_date=_to_finished_date(row[7] if len(row) > 7 else None),
                receipt_quantity=_to_decimal(row[8] if len(row) > 8 else None),
                notes=normalize_text(row[9] if len(row) > 9 else None),
                issue_quantity_1=_to_decimal(row[10] if len(row) > 10 else None),
                issue_date_1=_to_finished_date(row[11] if len(row) > 11 else None),
                issue_operator_1=normalize_text(row[12] if len(row) > 12 else None),
                issue_quantity_2=_to_decimal(row[13] if len(row) > 13 else None),
                issue_date_2=_to_finished_date(row[14] if len(row) > 14 else None),
                issue_operator_2=normalize_text(row[15] if len(row) > 15 else None),
                issue_quantity_3=_to_decimal(row[16] if len(row) > 16 else None),
                issue_date_3=_to_finished_date(row[17] if len(row) > 17 else None),
                issue_operator_3=normalize_text(row[18] if len(row) > 18 else None),
                project_code=normalize_text(row[19] if len(row) > 19 else None),
                producer_name=normalize_text(row[20] if len(row) > 20 else None),
                customer_name=normalize_text(row[21] if len(row) > 21 else None),
            )
        )

    return rows


def _build_finished_item(row: FinishedInventoryRow) -> FinishedInventoryItemRead:
    issue_dates = [row.issue_date_1, row.issue_date_2, row.issue_date_3]
    return FinishedInventoryItemRead(
        row_id=row.row_id,
        material_name=row.material_name,
        specification=row.specification,
        work_order_no=row.work_order_no,
        quantity_on_hand=row.quantity_on_hand,
        unit=row.unit,
        location_name=row.location_name,
        project_code=row.project_code,
        producer_name=row.producer_name,
        customer_name=row.customer_name,
        notes=row.notes,
        last_receipt_at=row.receipt_date,
        last_issue_at=max((value for value in issue_dates if value is not None), default=None),
    )


def _build_db_finished_item(item: FinishedInventoryItemModel) -> FinishedInventoryItemRead:
    return FinishedInventoryItemRead(
        row_id=_db_row_id(item.id),
        material_name=item.material_name,
        specification=item.specification,
        work_order_no=item.work_order_no,
        quantity_on_hand=item.quantity_on_hand,
        unit=item.unit,
        location_name=item.location_name,
        project_code=item.project_code,
        producer_name=item.producer_name,
        customer_name=item.customer_name,
        notes=item.notes,
        last_receipt_at=item.last_receipt_at,
        last_issue_at=item.last_issue_at,
    )


def _build_db_finished_transaction(transaction: FinishedInventoryTransactionModel) -> FinishedInventoryTransactionRead:
    return FinishedInventoryTransactionRead(
        id=f"db-{transaction.id}",
        row_id=_db_row_id(transaction.item_id),
        transaction_type=transaction.transaction_type.value,
        quantity=transaction.quantity,
        occurred_on=transaction.occurred_on,
        notes=transaction.notes,
    )


def create_finished_inventory_item(
    session: Session,
    *,
    material_name: str,
    specification: str | None,
    unit: str | None,
    location_name: str,
    quantity: Decimal,
    occurred_on: date,
    work_order_no: str | None = None,
    project_code: str | None = None,
    producer_name: str | None = None,
    customer_name: str | None = None,
    notes: str | None = None,
) -> tuple[FinishedInventoryItemRead, FinishedInventoryTransactionRead]:
    normalized_material_name = normalize_text(material_name)
    if not normalized_material_name:
        raise ValueError("Material name is required.")

    normalized_location_name = normalize_text(location_name)
    if not normalized_location_name:
        raise ValueError("Location is required.")

    normalized_specification = normalize_text(specification)
    normalized_unit = normalize_text(unit)
    normalized_work_order_no = normalize_text(work_order_no)
    normalized_project_code = normalize_text(project_code)
    normalized_producer_name = normalize_text(producer_name)
    normalized_customer_name = normalize_text(customer_name)
    normalized_notes = normalize_text(notes)

    excel_duplicate = next(
        (
            row
            for row in _load_finished_rows()
            if _same_finished_item(
                row,
                material_name=normalized_material_name,
                specification=normalized_specification,
                unit=normalized_unit,
                location_name=normalized_location_name,
            )
        ),
        None,
    )
    if excel_duplicate is not None:
        raise DuplicateFinishedInventoryItemError(excel_duplicate.row_id)

    db_duplicate = session.scalar(
        select(FinishedInventoryItemModel).where(
            and_(
                *_finished_match_conditions(
                    material_name=normalized_material_name,
                    specification=normalized_specification,
                    unit=normalized_unit,
                    location_name=normalized_location_name,
                )
            )
        )
    )
    if db_duplicate is not None:
        raise DuplicateFinishedInventoryItemError(_db_row_id(db_duplicate.id))

    normalized_quantity = Decimal(quantity)
    item = FinishedInventoryItemModel(
        material_name=normalized_material_name,
        specification=normalized_specification,
        work_order_no=normalized_work_order_no,
        quantity_on_hand=normalized_quantity,
        unit=normalized_unit,
        location_name=normalized_location_name,
        notes=normalized_notes,
        project_code=normalized_project_code,
        producer_name=normalized_producer_name,
        customer_name=normalized_customer_name,
        last_receipt_at=occurred_on,
    )
    session.add(item)
    session.flush()

    transaction = FinishedInventoryTransactionModel(
        item_id=item.id,
        transaction_type=TransactionType.receipt,
        quantity=normalized_quantity,
        occurred_on=occurred_on,
        notes=normalized_notes or "手动录入成品库存",
    )
    session.add(transaction)
    session.commit()
    session.refresh(item)
    session.refresh(transaction)

    return _build_db_finished_item(item), _build_db_finished_transaction(transaction)


def list_finished_inventory_items(session: Session | None = None) -> FinishedInventoryDashboardResponse:
    items = [_build_finished_item(row) for row in _load_finished_rows()]
    if session is not None:
        db_items = session.scalars(
            select(FinishedInventoryItemModel).order_by(
                FinishedInventoryItemModel.quantity_on_hand.asc(),
                FinishedInventoryItemModel.material_name.asc(),
            )
        ).all()
        items.extend(_build_db_finished_item(item) for item in db_items)

    total_stock_quantity = sum((item.quantity_on_hand for item in items), start=Decimal("0"))
    low_stock_items = sum(1 for item in items if item.quantity_on_hand <= LOW_STOCK_THRESHOLD)

    return FinishedInventoryDashboardResponse(
        summary=FinishedInventorySummary(
            total_items=len(items),
            total_stock_quantity=total_stock_quantity,
            low_stock_items=low_stock_items,
        ),
        items=items,
    )


def list_finished_inventory_transactions(row_id: int, session: Session | None = None) -> list[FinishedInventoryTransactionRead]:
    db_item_id = _db_item_id(row_id)
    if db_item_id is not None:
        if session is None:
            raise ValueError("Finished inventory row not found.")
        item = session.get(FinishedInventoryItemModel, db_item_id)
        if item is None:
            raise ValueError("Finished inventory row not found.")
        transactions = session.scalars(
            select(FinishedInventoryTransactionModel)
            .where(FinishedInventoryTransactionModel.item_id == db_item_id)
            .order_by(FinishedInventoryTransactionModel.occurred_on.desc(), FinishedInventoryTransactionModel.id.desc())
        ).all()
        return [_build_db_finished_transaction(transaction) for transaction in transactions]

    rows = _load_finished_rows()
    target = next((row for row in rows if row.row_id == row_id), None)
    if target is None:
        raise ValueError("Finished inventory row not found.")

    transactions: list[FinishedInventoryTransactionRead] = []

    if target.receipt_quantity > 0 or target.receipt_date is not None:
        transactions.append(
            FinishedInventoryTransactionRead(
                id=f"{row_id}-receipt",
                row_id=row_id,
                transaction_type="receipt",
                quantity=target.receipt_quantity,
                occurred_on=target.receipt_date,
                notes=target.notes or "由成品库存清单派生的最近入库记录",
            )
        )

    issue_slots = [
        (1, target.issue_quantity_1, target.issue_date_1, target.issue_operator_1),
        (2, target.issue_quantity_2, target.issue_date_2, target.issue_operator_2),
        (3, target.issue_quantity_3, target.issue_date_3, target.issue_operator_3),
    ]
    for slot, quantity, occurred_on, operator_name in issue_slots:
        if quantity <= 0 and occurred_on is None and operator_name is None:
            continue
        transactions.append(
            FinishedInventoryTransactionRead(
                id=f"{row_id}-issue-{slot}",
                row_id=row_id,
                transaction_type="issue",
                quantity=quantity,
                occurred_on=occurred_on,
                operator_name=operator_name,
                notes=f"由成品库存清单派生的出库记录 {slot}",
            )
        )

    transactions.sort(
        key=lambda item: (
            item.occurred_on or date.min,
            1 if item.transaction_type == "receipt" else 0,
            item.id,
        ),
        reverse=True,
    )
    return transactions
